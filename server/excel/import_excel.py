import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def split_tags(value):
    return [part.strip() for part in str(value or "").replace("，", "|").replace("、", "|").split("|") if part.strip()]


def row_dict(sheet, row):
    headers = [cell.value for cell in sheet[1]]
    return {headers[index]: row[index] for index in range(len(headers))}


def normalize_status(status):
    text = (status or "").strip()
    if "活动/联动" in text or "活动" in text or "联动" in text:
        return "活动/联动"
    if "更新" in text or "版本" in text:
        return "更新"
    return "上线/测试"


def infer_pool(game_type):
    if game_type == "福利":
        return "小七福利新游池"
    if game_type == "端游":
        return "PC/Steam池"
    return "公共手游池"


def build_source_record_map(wb):
    if "来源记录" not in wb.sheetnames:
        return {}
    sheet = wb["来源记录"]
    result = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        item = row_dict(sheet, row)
        game_id = str(item.get("游戏ID", "") or "").strip()
        if not game_id:
            continue
        result.setdefault(game_id, []).append({
            "source": str(item.get("来源", "") or "").strip(),
            "gameName": str(item.get("来源游戏名", "") or "").strip(),
            "status": normalize_status(str(item.get("状态", "") or "")),
            "eventType": normalize_status(str(item.get("事件", "") or "")),
            "eventAt": str(item.get("日期", "") or "").strip(),
            "tags": [normalize_status(str(item.get("状态", "") or "")), *split_tags(item.get("玩法标签", ""))],
            "vendor": str(item.get("游戏厂商", "") or "").strip(),
            "summary": str(item.get("简介", "") or "").strip(),
            "updateContent": str(item.get("补充内容", "") or "").strip(),
            "detailUrl": str(item.get("详情页", "") or "").strip(),
            "iconUrl": str(item.get("图标链接", "") or "").strip(),
        })
    return result


def build_mobile_games(wb):
    if "爬虫字段表" not in wb.sheetnames:
        raise ValueError("Excel 中找不到工作表：爬虫字段表")
    sheet = wb["爬虫字段表"]
    source_record_map = build_source_record_map(wb)
    rows = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        item = row_dict(sheet, row)
        enabled = str(item.get("启用前端", "") or "").strip()
        if enabled not in {"是", "yes", "YES", "true", "True", "1"}:
            continue
        game_id = str(item.get("游戏ID", "") or "").strip()
        if not game_id:
            continue
        game_type = str(item.get("游戏类型", "") or "手游").strip()
        status = normalize_status(str(item.get("开测类型", "") or "").strip())
        gameplay_tags = split_tags(item.get("玩法标签", ""))
        sources = split_tags(item.get("来源标签", ""))
        source_records = source_record_map.get(game_id, [])
        rows.append({
            "id": game_id,
            "gameName": str(item.get("游戏名称", "") or "").strip(),
            "subtitle": str(item.get("副标题", "") or "").strip(),
            "platform": game_type,
            "pool": infer_pool(game_type),
            "status": status,
            "eventType": status,
            "rawStatus": status,
            "rawEventType": status,
            "eventAt": source_records[0].get("eventAt", "") if source_records else "",
            "iconUrl": str(item.get("游戏图标", "") or "").strip(),
            "summary": str(item.get("简介", "") or "").strip(),
            "updateContent": str(item.get("补充内容", "") or "").strip(),
            "detailUrl": str(item.get("原文详情", "") or "").strip(),
            "originalUrl": str(item.get("原文详情", "") or "").strip(),
            "sources": sources,
            "gameplayTags": gameplay_tags,
            "tags": [status, *gameplay_tags],
            "sourceRecords": source_records,
            "mergedCount": max(len(source_records), len(sources), 1),
            "vendor": str(item.get("游戏厂商", "") or "").strip(),
        })
    return rows


def main():
    if len(sys.argv) < 4:
        raise SystemExit("usage: import_excel.py <xlsx_path> <source_json_path> <output_json_path>")
    xlsx_path = Path(sys.argv[1])
    source_json_path = Path(sys.argv[2])
    output_json_path = Path(sys.argv[3])
    source_data = json.loads(source_json_path.read_text(encoding="utf-8"))
    wb = load_workbook(xlsx_path)
    mobile_games = build_mobile_games(wb)
    source_data["mobileGames"] = mobile_games
    metrics = source_data.setdefault("metrics", {})
    metrics["totalGames"] = len(mobile_games) + len(source_data.get("welfareGames", [])) + len(source_data.get("pcGames", []))
    metrics["mobilePool"] = len(mobile_games)
    metrics["todayUpdates"] = len(mobile_games)
    output_json_path.write_text(json.dumps(source_data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(str(output_json_path))


if __name__ == "__main__":
    main()
