import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


FIELD_HEADERS = [
    "\u6e38\u620f\u56fe\u6807",
    "\u6e38\u620f\u540d\u79f0",
    "\u526f\u6807\u9898",
    "\u6e38\u620f\u7c7b\u578b",
    "\u73a9\u6cd5\u6807\u7b7e",
    "\u5f00\u6d4b\u7c7b\u578b",
    "\u6e38\u620f\u5382\u5546",
    "\u6e38\u620fID",
    "\u542f\u7528\u524d\u7aef",
    "\u6574\u7406\u5185\u5bb9",
    "\u539f\u6587\u8be6\u60c5",
    "\u6765\u6e90\u6807\u7b7e",
    "\u6765\u6e90\u8bb0\u5f55\u660e\u7ec6",
]

DATA_HEADERS = [
    "\u6765\u6e90Key",
    "\u6765\u6e90",
    "\u6e38\u620f\u540d\u79f0",
    "\u526f\u6807\u9898",
    "\u72b6\u6001",
    "\u4e8b\u4ef6",
    "\u65e5\u671f",
    "\u73a9\u6cd5\u6807\u7b7e",
    "\u6574\u7406\u5185\u5bb9",
    "\u6e38\u620f\u5382\u5546",
    "\u8be6\u60c5\u9875",
    "\u5217\u8868\u9875",
    "\u56fe\u6807\u94fe\u63a5",
    "\u95ee\u9898\u8bb0\u5f55",
]

SOURCE_RECORD_HEADERS = [
    "\u6e38\u620fID",
    "\u6e38\u620f\u540d\u79f0",
    "\u6765\u6e90",
    "\u6765\u6e90\u6e38\u620f\u540d",
    "\u72b6\u6001",
    "\u4e8b\u4ef6",
    "\u65e5\u671f",
    "\u73a9\u6cd5\u6807\u7b7e",
    "\u6e38\u620f\u5382\u5546",
    "\u6574\u7406\u5185\u5bb9",
    "\u8be6\u60c5\u9875",
    "\u56fe\u6807\u94fe\u63a5",
]

TAG_REVIEW_HEADERS = [
    "\u73a9\u6cd5\u6807\u7b7e",
    "\u51fa\u73b0\u6b21\u6570",
    "\u793a\u4f8b\u6e38\u620f",
    "\u6765\u6e90",
    "\u5f53\u524d\u72b6\u6001",
]


REVIEW_HEADERS = [
    "\u6e38\u620fID",
    "\u6e38\u620f\u540d\u79f0",
    "\u5408\u5e76\u540e\u6765\u6e90",
    "\u5ba1\u67e5\u5c42\u7ea7",
    "\u6765\u6e90",
    "\u6765\u6e90\u6e38\u620f\u540d",
    "\u5f53\u524d\u72b6\u6001",
    "\u5f53\u524d\u4e8b\u4ef6",
    "\u65e5\u671f",
    "\u6e38\u620f\u5382\u5546",
    "\u73a9\u6cd5\u6807\u7b7e",
    "\u7b80\u4ecb",
    "\u8865\u5145\u5185\u5bb9",
    "\u6574\u7406\u5185\u5bb9",
    "\u8be6\u60c5\u9875",
    "\u56fe\u6807\u94fe\u63a5",
    "\u8d28\u68c0\u95ee\u9898",
    "\u7528\u6237\u5ba1\u67e5\u610f\u89c1",
    "\u7528\u6237\u6307\u5b9a\u4fdd\u7559\u5b57\u6bb5",
]

README_ROWS = [
    ["\u5b57\u6bb5", "\u8bf4\u660e"],
    ["\u6e38\u620f\u56fe\u6807", "\u5f53\u524d\u53ef\u7528\u7684\u6e38\u620f\u56fe\u6807\u94fe\u63a5\uff0c\u524d\u7aef\u9884\u7559 ICON \u4f4d\u4f1a\u4f7f\u7528\u5b83\u3002"],
    ["\u6e38\u620f\u540d\u79f0", "\u6e05\u6d17\u540e\u7684\u4e3b\u6e38\u620f\u540d\uff0c\u4e0d\u5305\u542b\u526f\u6807\u9898\u3002"],
    ["\u526f\u6807\u9898", "\u4ece\u77ed\u6a2a\u7ebf\u6216\u62ec\u53f7\u4e2d\u62c6\u51fa\u7684\u8bf4\u660e\u6587\u6848\u3002"],
    ["\u6e38\u620f\u7c7b\u578b", "\u624b\u6e38 / \u798f\u5229 / \u7aef\u6e38\uff0c\u6309\u6293\u53d6\u6765\u6e90\u6c60\u5212\u5206\u3002"],
    ["\u73a9\u6cd5\u6807\u7b7e", "\u6e38\u620f\u672c\u8eab\u73a9\u6cd5\u6807\u7b7e\uff0c\u5df2\u6392\u9664\u7f16\u8f91\u63a8\u8350\u3001\u9ad8\u753b\u8d28\u3001\u4e70\u65ad\u5236\u7b49\u975e\u73a9\u6cd5\u6807\u7b7e\u3002"],
    ["\u5f00\u6d4b\u7c7b\u578b", "\u7edf\u4e00\u5408\u5e76\u4e3a\uff1a\u4e0a\u7ebf/\u6d4b\u8bd5\u3001\u66f4\u65b0\u3001\u6d3b\u52a8/\u8054\u52a8\u3002"],
    ["\u6e38\u620f\u5382\u5546", "\u4ece\u6765\u6e90\u8be6\u60c5\u9875\u4f18\u5148\u8865\u5145\u3002"],
    ["\u6574\u7406\u5185\u5bb9", "\u5df2\u5408\u5e76\u7b80\u4ecb\u4e0e\u8865\u5145\u5185\u5bb9\uff0c\u4f18\u5148\u4fdd\u7559\u5f00\u6d4b\u3001\u5f00\u670d\u3001\u6d3b\u52a8\u3001\u7248\u672c\u66f4\u65b0\u7b49\u6709\u6548\u53e5\u5b50\u3002"],
    ["\u539f\u6587\u8be6\u60c5", "\u6765\u6e90\u7f51\u7ad9\u5bf9\u5e94\u6e38\u620f\u8be6\u60c5\u9875\u3002"],
    ["\u6765\u6e90\u8bb0\u5f55\u660e\u7ec6", "\u805a\u5408\u6e38\u620f\u4e0b\u9762\u7684\u5404\u5e73\u53f0\u539f\u59cb\u6765\u6e90\u660e\u7ec6\u3002"],
]

TAG_BLOCKLIST = {
    "\u4e0a\u7ebf/\u6d4b\u8bd5", "\u66f4\u65b0", "\u6d3b\u52a8/\u8054\u52a8", "\u65b0\u6e38", "\u5f00\u6d4b", "\u9884\u7ea6", "\u9996\u53d1", "\u4e0a\u7ebf",
    "\u4eca\u65e5\u4e0a\u7ebf", "\u5373\u5c06\u4e0a\u7ebf", "\u7248\u672c\u66f4\u65b0", "\u6d3b\u52a8", "\u8054\u52a8", "\u52a8\u6001", "\u4e0b\u8f7d", "\u8bd5\u73a9",
    "\u5173\u6ce8", "\u653b\u7565", "\u4eae\u70b9", "\u8be6\u60c5", "[\u8be6\u60c5]", "\u65b0\u624b", "\u65b0\u95fb", "\u6e38\u620f", "\u5f00\u6d4b\u63d0\u9192",
    "\u9884\u7ea6\u63d0\u9192", "\u7f16\u8f91\u63a8\u8350", "\u529b\u8350\u4f73\u4f5c", "\u5fc5\u73a9\u795e\u4f5c", "\u7cbe\u54c1", "\u4f18\u9009", "\u9ad8\u753b\u8d28", "\u4e70\u65ad\u5236",
    "\u6c89\u6d78\u5f0f\u4f53\u9a8c", "\u6a2a\u7248", "\u8d85\u723d\u6253\u51fb", "\u591a\u5e73\u53f0", "\u5b98\u670d", "\u5b89\u5353", "iOS", "PC", "\u6d88\u901d\u7684\u6e38\u620f", "\u9996\u9875", "\u6392\u884c\u699c", "\u699c\u5355", "\u4e13\u533a", "\u8bba\u575b", "\u793e\u533a", "\u5b98\u65b9", "\u5b98\u7f51",
    "\u4e91\u6e38\u620f", "\u81ea\u52a8\u62fe\u53d6", "\u81ea\u52a8", "\u5212\u5206", "\u4e0a\u5343\u5730\u56fe\u7b49\u4f60\u63a2\u7d22",
    "\u7b80\u4ecb", "\u8ba1\u8d39\u5220\u6863\u5185\u6d4b", "\u4ed8\u8d39\u4e0b\u8f7d", "\u591a\u7ed3\u5c40", "\u53d1\u884c\u5546", "\u5c81\u672b\u5e86\u5178\u5f00\u542f", "\u70ed\u95e8\u6807\u7b7e\uff1a", "\u810d\u7099\u4eba\u53e3", "CV\u63a7", "up\u4e3b\u63a8\u8350", "\u4efb\u52a1", "\u526f\u672c", "\u52a0\u901f\u5668\u4e0b\u8f7d", "\u53d1\u5e03", "\u65e0\u53cc\u9a6c\u8d85\u767b\u573a\uff01", "\u79c1\u85cf\u4f73\u4f5c", "\u67e5\u770b\u66f4\u591a>>", "BOSS\u8ba8\u4f10", "\u4e0d\u5220\u6863\u5185\u6d4b", "\u4e3b\u6253\u590d\u53e4", "\u4e5d\u6e38APP", "\u4ecb\u7ecd\u4e4b\u516c\u4f1a\u6218", "\u4e92\u52a8\u5f71\u50cf", "\u4f60\u753b\u6211\u731c", "\u4f60\u7684\u4f20\u5947\u7531\u4f60\u5b9a\u4e49",
    "\u65b0\u73a9\u6cd5", "\u65b0\u5730\u56fe", "\u65b0\u5730", "\u804c\u4e1a", "\u5e03\u9635\u8fce\u654c", "\u5b88\u62a4\u57ce\u5821", "\u62b5\u5fa1\u4e00",
}

TAG_BLOCK_PATTERNS = [
    re.compile(r"(\u7248\u672c|\u66f4\u65b0|\u6d3b\u52a8|\u8d5b\u5b63|\u8054\u52a8|\u8fd4\u573a|\u4e0a\u7ebf|\u9884\u7ea6|\u6d4b\u8bd5|\u9996\u6d4b|\u7ec8\u6d4b|\u9996\u53d1|\u63d0\u9192|\u5220\u6863|\u5185\u6d4b|\u516c\u6d4b|\u4e0b\u8f7d|\u53d1\u884c|\u5382\u5546|\u7b80\u4ecb|BOSS)"),
    re.compile(r"^(\u7c7b\u578b|\u6807\u7b7e|\u5206\u7c7b|\u73a9\u6cd5|\u5f00\u53d1|\u66f4\u591a|\u8be6\u60c5|\u8ba2\u9605|\[\u8be6\u60c5\])$"),
]

GAMEPLAY_TAG_ALLOWLIST = {
    "角色扮演", "动作", "策略", "战棋", "卡牌", "射击", "竞技", "多人联机", "开放世界", "二次元",
    "养成", "生存", "像素", "模拟", "模拟经营", "休闲", "格斗", "沙盒", "解谜", "音游",
    "社交", "AVG", "MMO", "MOBA", "SLG", "ARPG", "JRPG", "TPS", "TCG", "DBG", "3D",
    "Roguelite", "RogueLite", "Roguelike", "RogueLike", "Steam移植",
    "竞速", "赛车", "篮球", "体育", "足球", "塔防", "冒险", "刷宝", "抓宠", "自走棋",
    "非对称竞技", "派对游戏", "搜打撤", "搜打拆", "枪战", "弹幕", "涂色", "建造", "割草", "剪草",
    "剧情", "回合", "即时战斗", "即时策略", "经营", "收集", "暗黑", "地牢", "女性向", "恋爱",
    "乙女", "换装", "文字", "视觉小说", "推理", "悬疑", "恐怖惊悚", "末日", "科幻", "魔幻",
    "武侠", "修仙", "仙侠", "中国风", "都市", "克苏鲁", "农场模拟", "密室逃脱", "寻物", "平台跳跃",
    "消除", "海战", "街机", "跑酷", "迷宫", "狩猎", "钓鱼", "美食", "益智", "跑团",
    "冒险解谜", "卡通", "温馨治愈", "手绘", "唯美", "小清新", "末日生存", "装饰&装修",
    "城市建造", "农场", "种田", "漂移", "美少女", "偶像", "学园", "动漫改编", "小说改编",
    "页游改编", "端游改编", "类逃离塔科夫", "类银河恶魔城", "狼人杀", "吃鸡", "魂like",
}

TAG_ALIAS = {
    "搜打拆": "搜打撤",
    "下落式音游": "音游",
    "节奏": "音游",
    "飞行射击": "射击",
    "城市建造": "建造",
    "农场": "农场模拟",
    "RogueLite": "Roguelite",
    "RogueLike": "Roguelike",
}
CONTENT_BLOCK_PATTERNS = [
    re.compile(r"\{\s*['\"]?@type"),
    re.compile(r"privacy_policy|acceptedAnswer|application/ld\+json|t\d+_\d+-t\d+_\d+|apk_size|version_name", re.I),
    re.compile(r"\u597d\u6e38\u5feb\u7206APP\u63d0\u4f9b|\u4e5d\u6e38.*\u63d0\u4f9b.*(\u9884\u7ea6|\u4e0b\u8f7d|\u653b\u7565)"),
    re.compile(r"\u6765\u81ea\u4e8e\s*\u4e5d\u6e38APP|\u8c4c\u8c46"),
]


def stringify(value):
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        return " | ".join(str(item) for item in value if item is not None and str(item) != "")
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def split_tag_text(value):
    return [part.strip() for part in str(value or "").replace("\uff0c", "|").replace("\u3001", "|").replace(",", "|").split("|") if part.strip()]


def normalize_tag_alias(tag):
    tag = str(tag or "").strip().replace("搜打拆", "搜打撤")
    return TAG_ALIAS.get(tag, tag)


def clean_gameplay_tags(values):
    tags = []
    for value in values or []:
        for raw_tag in split_tag_text(value):
            tag = normalize_tag_alias(raw_tag)
            if len(tag) < 2 or len(tag) > 12:
                continue
            if tag in TAG_BLOCKLIST:
                continue
            if re.match(r"^\d", tag) or re.search(r"\d{1,2}[:\-]\d{2}|20\d{2}|^\(\d+\)$|^\+\d+", tag):
                continue
            if re.search(r"[。！？!?，,《》“”\"']", tag):
                continue
            if any(pattern.search(tag) for pattern in TAG_BLOCK_PATTERNS):
                continue
            if tag not in GAMEPLAY_TAG_ALLOWLIST:
                continue
            if tag not in tags:
                tags.append(tag)
    return tags[:12]

def trim_update_history_text(value):
    text = str(value or "").strip()
    matches = list(re.finditer(r"20\d{2}[.\/-]\d{1,2}[.\/-]\d{1,2}", text))
    if len(matches) >= 2:
        return text[:matches[1].start()].strip()
    if len(matches) == 1 and matches[0].start() > 12:
        return text[:matches[0].start()].strip()
    return text


def clean_content_piece(value):
    text = str(value or "").strip()
    if not text:
        return ""
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"\u597d\u6e38\u5feb\u7206APP\u63d0\u4f9b[^\u300a]{0,80}(\u9884\u7ea6|\u4e0b\u8f7d)?", "", text)
    text = re.sub(r"\u4e5d\u6e38[^\u3002]{0,80}\u63d0\u4f9b[^\u3002]{0,80}(\u9884\u7ea6|\u4e0b\u8f7d|\u653b\u7565)", "", text)
    text = re.sub(r"t\d+_\d+-t\d+_\d+[:\uff1a]\d+(?:\.\d+)?", "", text)
    text = text.replace("-->", " ").replace("[\u8be6\u60c5]", "")
    text = re.sub(r"[\uff0c,]?TapTap\s*\u63d0\u4f9b[^\uff0c\u3002]{0,100}[\uff0c\u3002]?", " ", text)
    text = text.replace("\u6700\u65b0\u7248\u5b98\u65b9\u6b63\u7248", " ")
    text = re.sub(r"\s+", " ", text).strip()
    if any(pattern.search(text) for pattern in CONTENT_BLOCK_PATTERNS):
        return ""
    if any(word in text for word in ["\u7248\u672c\u8981\u6c42", "\u65e0\u7cfb\u7edf\u8981\u6c42", "\u5f53\u524d\u7248\u672c", "\u66f4\u65b0\u65f6\u95f4"]):
        return ""
    text = trim_update_history_text(text)
    if text in {"\u9884\u7ea6", "\u9996\u53d1", "\u4e0a\u7ebf", "\u5f00\u6d4b", "\u52a8\u6001", "\u6d3b\u52a8", "\u66f4\u65b0", "\u7248\u672c\u66f4\u65b0", "\u4e0b\u8f7d", "\u5173\u6ce8", "\u641c\u7d22"}:
        return ""
    return text


def similar_content(left, right):
    def norm(value):
        return re.sub(r"[\s\uff0c\u3002\uff01\uff1f\u3001,.!?:\uff1a\u300a\u300b\u201c\u201d'\"\uff08\uff09()]", "", str(value or "")).lower()
    a, b = norm(left), norm(right)
    if not a or not b:
        return False
    return a == b or (len(a) >= 10 and a in b) or (len(b) >= 10 and b in a)


def merged_content(*values):
    pieces = []
    for value in values:
        for part in re.split(r"\n+|(?<=[\u3002\uff01\uff1f])", str(value or "")):
            cleaned = clean_content_piece(part)
            if cleaned and not any(similar_content(cleaned, existing) for existing in pieces):
                pieces.append(cleaned)
    return "\n".join(pieces[:2])


SOURCE_PRIORITY = ["\u597d\u6e38\u5feb\u7206", "TapTap", "\u4e5d\u6e38"]


def single_content(*values):
    text = merged_content(*values)
    return next((line.strip() for line in text.splitlines() if line.strip()), "")


def source_rank(source):
    source_text = str(source or "")
    for index, name in enumerate(SOURCE_PRIORITY):
        if source_text == name:
            return index
    return len(SOURCE_PRIORITY)


def priority_content(item):
    records = sorted(item.get("sourceRecords", []) or [], key=lambda record: source_rank(record.get("source")))
    for record in records:
        content = single_content(record.get("summary", ""), record.get("updateContent", ""))
        if content:
            return content
    return single_content(item.get("summary", ""), item.get("updateContent", ""))


def gameplay_tags(item):
    blocked = {item.get("status"), item.get("eventType"), item.get("rawStatus"), item.get("rawEventType"), *(item.get("sources") or [])}
    values = item.get("gameplayTags") or item.get("tags") or []
    return [tag for tag in clean_gameplay_tags(values) if tag not in blocked]


def source_records_text(records):
    lines = []
    for record in records or []:
        parts = [
            stringify(record.get("source")),
            stringify(record.get("status")),
            stringify(record.get("eventType")),
            stringify(record.get("eventAt")),
            stringify(record.get("vendor")),
            stringify(record.get("detailUrl")),
        ]
        lines.append(" / ".join(part for part in parts if part))
    return "\n".join(lines)


def issues_text(issues, title):
    related = [issue for issue in issues or [] if issue.get("title") == title]
    return "\n".join(f'{issue.get("type", "")}: {issue.get("detail", "")}' for issue in related)


def diagnose_text_quality(value=""):
    text = str(value or "").strip()
    reasons = []
    if not text:
        reasons.append("\u7f3a\u5185\u5bb9")
    elif len(text) < 8:
        reasons.append("\u5185\u5bb9\u8fc7\u77ed")
    if len(text) > 240:
        reasons.append("\u5185\u5bb9\u8fc7\u957f")
    noise_words = [
        "\u4e0b\u8f7d", "\u653b\u7565", "\u793c\u5305", "\u4e13\u533a", "\u4e5d\u6e38APP", "\u597d\u6e38\u5feb\u7206APP", "\u9690\u79c1", "\u7528\u6237\u534f\u8bae",
        "\u9884\u7ea6\u4e0b\u8f7d", "\u52a0\u901f\u5668", "\u67e5\u770b\u66f4\u591a", "\u70ed\u95e8\u6807\u7b7e", "\u8c4c\u8c46\u835a", "\u4fdd\u7559\u6e38\u620f\u4e13\u533a\u8be6\u60c5\u9875",
    ]
    hits = [word for word in noise_words if word in text]
    if hits:
        reasons.append("\u542b\u9875\u9762\u566a\u97f3\uff1a" + " / ".join(hits[:3]))
    if re.fullmatch(r"(\u63d0\u9192|\u9884\u7ea6|\u9996\u53d1|\u4e0a\u7ebf|\u5f00\u6d4b|\u65b0\u6e38|\u4e0b\u8f7d|\u52a8\u6001|\u641c\u7d22|\u8be6\u60c5)", text):
        reasons.append("\u7591\u4f3c\u6309\u94ae\u6216\u72b6\u6001\u6587\u6848")
    if re.search(r"[a-zA-Z0-9_]{16,}", text):
        reasons.append("\u7591\u4f3c\u4ee3\u7801\u6216\u94fe\u63a5\u6b8b\u7559")
    if any(pattern.search(text) for pattern in CONTENT_BLOCK_PATTERNS):
        reasons.append("\u7591\u4f3c\u7ed3\u6784\u5316\u9875\u9762\u6b8b\u7559")
    return reasons


def diagnose_tag_quality(tags=None):
    values = tags if isinstance(tags, list) else split_tag_text(tags or "")
    reasons = []
    if not values:
        reasons.append("\u7f3a\u73a9\u6cd5\u6807\u7b7e")
    noisy = []
    for tag in values:
        tag_text = str(tag or "").strip()
        if not tag_text:
            continue
        if tag_text in TAG_BLOCKLIST or any(pattern.search(tag_text) for pattern in TAG_BLOCK_PATTERNS):
            noisy.append(tag_text)
        elif re.search(r"\u4e0a\u6709\u5927\u91cf|\u6309\u6807\u7b7e\u6d4f\u89c8|\u60a8\u53ef\u4ee5\u5728|\u70ed\u95e8\u6807\u7b7e|\u7f16\u8f91\u63a8\u8350|\u4e0b\u8f7d|APP|\u4e13\u533a|\u67e5\u770b\u66f4\u591a", tag_text):
            noisy.append(tag_text)
        elif len(tag_text) > 14:
            noisy.append(tag_text)
    if noisy:
        reasons.append("\u6807\u7b7e\u542b\u566a\u97f3\uff1a" + " / ".join(noisy[:3]))
    return reasons


def diagnose_review_row(record, tags):
    reasons = []
    reasons.extend(diagnose_text_quality(record.get("summary", "") or record.get("updateContent", "")))
    reasons.extend(diagnose_tag_quality(tags))
    if not record.get("vendor"):
        reasons.append("\u7f3a\u5382\u5546")
    if not record.get("detailUrl"):
        reasons.append("\u7f3a\u8be6\u60c5\u9875")
    if not record.get("iconUrl"):
        reasons.append("\u7f3aICON")
    return "\uff1b".join(dict.fromkeys(reasons))


def autosize(ws):
    for column in ws.columns:
        length = 0
        letter = column[0].column_letter
        for cell in column:
            value = "" if cell.value is None else str(cell.value)
            length = max(length, len(value))
        ws.column_dimensions[letter].width = min(max(length + 2, 10), 54)


def style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="21A1F1")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def normalize_game_type(item):
    pool = str(item.get("pool", "") or "")
    if "\u798f\u5229" in pool:
        return "\u798f\u5229"
    if "PC" in pool or "Steam" in pool or "\u7aef\u6e38" in pool:
        return "\u7aef\u6e38"
    return "\u624b\u6e38"


def normalize_event_type(value):
    text = str(value or "")
    if "\u6d3b\u52a8/\u8054\u52a8" in text or "\u6d3b\u52a8" in text or "\u8054\u52a8" in text:
        return "\u6d3b\u52a8/\u8054\u52a8"
    if "\u66f4\u65b0" in text or "\u7248\u672c" in text:
        return "\u66f4\u65b0"
    return "\u4e0a\u7ebf/\u6d4b\u8bd5"


def build_field_sheet(wb, data):
    ws = wb.active
    ws.title = "\u722c\u866b\u5b57\u6bb5\u8868"
    ws.append(FIELD_HEADERS)
    for item in data.get("mobileGames", []):
        ws.append([
            item.get("iconUrl", ""),
            item.get("gameName", ""),
            item.get("subtitle", "") or "",
            normalize_game_type(item),
            stringify(gameplay_tags(item)),
            normalize_event_type(item.get("status", "") or item.get("eventType", "")),
            item.get("vendor", "") or "",
            item.get("id", ""),
            "\u662f",
            priority_content(item),
            item.get("detailUrl", ""),
            stringify(item.get("sources", [])),
            source_records_text(item.get("sourceRecords", [])),
        ])
    style_sheet(ws)
    autosize(ws)


def build_data_sheet(wb, data):
    ws = wb.create_sheet("\u722c\u53d6\u6570\u636e\u8868")
    ws.append(DATA_HEADERS)
    for item in data.get("rawMobileRecords", []):
        ws.append([
            item.get("sourceKey", ""),
            item.get("source", ""),
            item.get("gameName", ""),
            item.get("subtitle", "") or "",
            item.get("status", ""),
            item.get("eventType", ""),
            item.get("eventAt", ""),
            stringify(gameplay_tags(item)),
            single_content(item.get("summary", ""), item.get("updateContent", "")),
            item.get("vendor", "") or "",
            item.get("detailUrl", ""),
            item.get("originalUrl", ""),
            item.get("iconUrl", ""),
            issues_text(data.get("cleanupIssues", []), item.get("gameName", "")),
        ])
    style_sheet(ws)
    autosize(ws)


def build_source_record_sheet(wb, data):
    ws = wb.create_sheet("\u6765\u6e90\u8bb0\u5f55")
    ws.append(SOURCE_RECORD_HEADERS)
    for item in data.get("mobileGames", []):
        for record in item.get("sourceRecords", []) or []:
            ws.append([
                item.get("id", ""),
                item.get("gameName", ""),
                record.get("source", ""),
                record.get("gameName", ""),
                record.get("status", ""),
                record.get("eventType", ""),
                record.get("eventAt", ""),
                stringify(gameplay_tags(record)),
                record.get("vendor", "") or "",
                single_content(record.get("summary", ""), record.get("updateContent", "")),
                record.get("detailUrl", ""),
                record.get("iconUrl", ""),
            ])
    style_sheet(ws)
    autosize(ws)





def load_review_overrides(xlsx_path):
    path = Path(xlsx_path)
    if not path.exists():
        return {}
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        if "爬虫字段审查表" not in wb.sheetnames:
            wb.close()
            return {}
        ws = wb["爬虫字段审查表"]
        headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
        index = {name: pos + 1 for pos, name in enumerate(headers) if name}
        required = ["游戏名称", "审查层级", "来源", "来源游戏名", "详情页"]
        if not all(name in index for name in required):
            wb.close()
            return {}
        overrides = {}
        for row in range(2, ws.max_row + 1):
            key = tuple(str(ws.cell(row, index[name]).value or "") for name in required)
            opinion = ws.cell(row, index.get("用户审查意见", 0)).value if index.get("用户审查意见") else ""
            keep = ws.cell(row, index.get("用户指定保留字段", 0)).value if index.get("用户指定保留字段") else ""
            if opinion or keep:
                overrides[key] = (opinion or "", keep or "")
        wb.close()
        return overrides
    except Exception:
        return {}


def review_override(overrides, game_name, level, source, source_game_name, detail_url):
    return overrides.get(tuple(str(value or "") for value in [game_name, level, source, source_game_name, detail_url]), ("", ""))

def build_review_sheet(wb, data, review_overrides=None):
    ws = wb.create_sheet("\u722c\u866b\u5b57\u6bb5\u5ba1\u67e5\u8868")
    ws.append(REVIEW_HEADERS)
    review_overrides = review_overrides or {}
    for item in data.get("mobileGames", []):
        sources = stringify(item.get("sources", []))
        cleaned_tags = gameplay_tags(item)
        cleaned_summary = priority_content(item)
        opinion, keep = review_override(review_overrides, item.get("gameName", ""), "\u6e05\u6d17\u540e", "\u5408\u5e76\u7ed3\u679c", item.get("gameName", ""), item.get("detailUrl", ""))
        ws.append([
            item.get("id", ""),
            item.get("gameName", ""),
            sources,
            "\u6e05\u6d17\u540e",
            "\u5408\u5e76\u7ed3\u679c",
            item.get("gameName", ""),
            item.get("status", ""),
            item.get("eventType", ""),
            item.get("eventAt", ""),
            item.get("vendor", "") or "",
            stringify(cleaned_tags),
            item.get("summary", "") or "",
            item.get("updateContent", "") or "",
            cleaned_summary,
            item.get("detailUrl", ""),
            item.get("iconUrl", ""),
            diagnose_review_row({**item, "summary": cleaned_summary}, cleaned_tags),
            opinion,
            keep,
        ])
        for record in item.get("sourceRecords", []) or []:
            raw_tags = record.get("gameplayTags") or record.get("tags") or []
            row_summary = single_content(record.get("summary", ""), record.get("updateContent", ""))
            opinion, keep = review_override(review_overrides, item.get("gameName", ""), "\u6765\u6e90\u539f\u59cb", record.get("source", ""), record.get("gameName", ""), record.get("detailUrl", ""))
            ws.append([
                item.get("id", ""),
                item.get("gameName", ""),
                sources,
                "\u6765\u6e90\u539f\u59cb",
                record.get("source", ""),
                record.get("gameName", ""),
                record.get("status", ""),
                record.get("eventType", ""),
                record.get("eventAt", ""),
                record.get("vendor", "") or "",
                stringify(raw_tags),
                record.get("summary", "") or "",
                record.get("updateContent", "") or "",
                row_summary,
                record.get("detailUrl", ""),
                record.get("iconUrl", ""),
                diagnose_review_row(record, raw_tags),
                opinion,
                keep,
            ])
    style_sheet(ws)
    autosize(ws)


def build_tag_review_sheet(wb, data):
    ws = wb.create_sheet("\u73a9\u6cd5\u6807\u7b7e\u6e05\u5355")
    ws.append(TAG_REVIEW_HEADERS)
    stats = {}

    def add_tag(tag, game_name, source, status):
        tag = str(tag or "").strip()
        if not tag:
            return
        bucket = stats.setdefault(tag, {"count": 0, "games": [], "sources": set(), "statuses": set()})
        bucket["count"] += 1
        if game_name and game_name not in bucket["games"]:
            bucket["games"].append(game_name)
        if source:
            bucket["sources"].add(source)
        if status:
            bucket["statuses"].add(status)

    for item in data.get("mobileGames", []):
        game_name = item.get("gameName", "")
        kept = set(gameplay_tags(item))
        raw_values = list(item.get("gameplayTags") or []) + list(item.get("tags") or [])
        for record in item.get("sourceRecords", []) or []:
            raw_values.extend(record.get("tags") or [])
        for raw in raw_values:
            for tag in split_tag_text(raw):
                status = "\u4fdd\u7559" if tag in kept or tag in clean_gameplay_tags([tag]) else "\u5df2\u8fc7\u6ee4"
                add_tag(tag, game_name, "\u539f\u59cb/\u5408\u5e76", status)

    for record in data.get("rawMobileRecords", []) or []:
        game_name = record.get("gameName", "")
        source = record.get("source", "") or "\u539f\u59cb\u8bb0\u5f55"
        for raw in record.get("tags", []) or []:
            for tag in split_tag_text(raw):
                status = "\u4fdd\u7559" if clean_gameplay_tags([tag]) else "\u5df2\u8fc7\u6ee4"
                add_tag(tag, game_name, source, status)

    for tag, info in sorted(stats.items(), key=lambda pair: ("\u5df2\u8fc7\u6ee4" in pair[1]["statuses"], -pair[1]["count"], pair[0])):
        ws.append([
            tag,
            info["count"],
            " | ".join(info["games"][:8]),
            " | ".join(sorted(info["sources"])),
            " | ".join(sorted(info["statuses"])),
        ])
    style_sheet(ws)
    autosize(ws)


def build_readme_sheet(wb):
    ws = wb.create_sheet("\u8bf4\u660e")
    for row in README_ROWS:
        ws.append(row)
    style_sheet(ws)
    autosize(ws)


def main():
    if len(sys.argv) < 3:
        raise SystemExit("usage: export_excel.py <json_path> <xlsx_path>")
    json_path = Path(sys.argv[1])
    xlsx_path = Path(sys.argv[2])
    data = json.loads(json_path.read_text(encoding="utf-8"))
    review_overrides = load_review_overrides(xlsx_path)
    wb = Workbook()
    build_field_sheet(wb, data)
    build_data_sheet(wb, data)
    build_source_record_sheet(wb, data)
    build_review_sheet(wb, data, review_overrides)
    build_tag_review_sheet(wb, data)
    build_readme_sheet(wb)
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)
    print(str(xlsx_path))


if __name__ == "__main__":
    main()



