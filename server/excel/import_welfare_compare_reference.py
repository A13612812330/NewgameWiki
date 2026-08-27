import argparse
import hashlib
import json
import re
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


NAME_KEYS = {"游戏名称", "游戏名", "名称", "name", "gameName"}
SUBTITLE_KEYS = {"副标题", "折扣标签", "福利标签", "版本", "subtitle", "subTitle"}
CATEGORY_KEYS = {"类别", "分类", "类型", "category"}


def now_iso():
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def normalize_header(value):
    return str(value or "").strip()


def normalize_text(value):
    return re.sub(
        r"[\s\"'“”‘’`~!！?？,，.。:：;；、|/\\\-_=+·・《》<>【】\[\]（）(){}]",
        "",
        str(value or ""),
    ).lower()


def find_column(headers, candidates):
    normalized = [normalize_header(item) for item in headers]
    for candidate in candidates:
        if candidate in normalized:
            return normalized.index(candidate)
    for index, header in enumerate(normalized):
        if any(candidate in header for candidate in candidates if len(candidate) >= 2):
            return index
    return -1


def parse_workbook(xlsx_path):
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
      records = []
      first_sheet_name = workbook.sheetnames[0] if workbook.sheetnames else ""
      for sheet in workbook.worksheets:
          rows = sheet.iter_rows(values_only=True)
          headers = next(rows, None)
          if not headers:
              continue
          name_index = find_column(headers, NAME_KEYS)
          subtitle_index = find_column(headers, SUBTITLE_KEYS)
          category_index = find_column(headers, CATEGORY_KEYS)
          if name_index < 0:
              continue
          for row_number, row in enumerate(rows, start=2):
              game_name = str(row[name_index] or "").strip() if name_index < len(row) else ""
              if not game_name:
                  continue
              subtitle = str(row[subtitle_index] or "").strip() if 0 <= subtitle_index < len(row) else ""
              category = str(row[category_index] or "").strip() if 0 <= category_index < len(row) else ""
              name_key = normalize_text(game_name)
              subtitle_key = normalize_text(subtitle)
              records.append({
                  "sheet": sheet.title,
                  "rowNumber": row_number,
                  "gameName": game_name,
                  "subtitle": subtitle,
                  "category": category,
                  "nameKey": name_key,
                  "subtitleKey": subtitle_key,
                  "matchKey": f"{name_key}::{subtitle_key}",
              })
      return first_sheet_name, records
    finally:
        workbook.close()


def ensure_schema(connection):
    connection.executescript(
        """
        CREATE TABLE IF NOT EXISTS welfare_compare_reference_imports (
          id TEXT PRIMARY KEY,
          file_name TEXT NOT NULL,
          source_path TEXT,
          file_sha256 TEXT NOT NULL UNIQUE,
          file_size INTEGER NOT NULL DEFAULT 0,
          sheet_name TEXT,
          row_count INTEGER NOT NULL DEFAULT 0,
          imported_at TEXT NOT NULL,
          is_active INTEGER NOT NULL DEFAULT 1,
          file_blob BLOB
        );
        CREATE TABLE IF NOT EXISTS welfare_compare_reference_records (
          id TEXT PRIMARY KEY,
          import_id TEXT NOT NULL,
          row_number INTEGER NOT NULL,
          game_name TEXT NOT NULL,
          subtitle TEXT,
          category TEXT,
          name_key TEXT NOT NULL,
          subtitle_key TEXT NOT NULL DEFAULT '',
          match_key TEXT NOT NULL,
          raw_json TEXT DEFAULT '{}',
          created_at TEXT NOT NULL,
          FOREIGN KEY(import_id) REFERENCES welfare_compare_reference_imports(id)
        );
        CREATE TABLE IF NOT EXISTS operation_logs (
          id TEXT PRIMARY KEY,
          type TEXT NOT NULL,
          target_type TEXT,
          target_id TEXT,
          title TEXT,
          detail TEXT,
          payload_json TEXT DEFAULT '{}',
          created_at TEXT NOT NULL
        );
        CREATE INDEX IF NOT EXISTS idx_welfare_compare_reference_active ON welfare_compare_reference_imports(is_active, imported_at);
        CREATE INDEX IF NOT EXISTS idx_welfare_compare_reference_match ON welfare_compare_reference_records(match_key);
        CREATE INDEX IF NOT EXISTS idx_welfare_compare_reference_name ON welfare_compare_reference_records(name_key);
        """
    )


def import_reference(db_path, xlsx_path):
    file_bytes = xlsx_path.read_bytes()
    file_hash = hashlib.sha256(file_bytes).hexdigest()
    import_id = f"welfare-ref-{file_hash[:16]}"
    imported_at = now_iso()
    sheet_name, records = parse_workbook(xlsx_path)
    connection = sqlite3.connect(db_path)
    try:
        ensure_schema(connection)
        connection.execute("BEGIN")
        connection.execute("UPDATE welfare_compare_reference_imports SET is_active = 0")
        connection.execute(
            """
            INSERT INTO welfare_compare_reference_imports
              (id, file_name, source_path, file_sha256, file_size, sheet_name, row_count, imported_at, is_active, file_blob)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, 1, ?)
            ON CONFLICT(file_sha256) DO UPDATE SET
              file_name = excluded.file_name,
              source_path = excluded.source_path,
              file_size = excluded.file_size,
              sheet_name = excluded.sheet_name,
              row_count = excluded.row_count,
              imported_at = excluded.imported_at,
              is_active = 1,
              file_blob = excluded.file_blob
            """,
            (import_id, xlsx_path.name, str(xlsx_path), file_hash, len(file_bytes), sheet_name, len(records), imported_at, sqlite3.Binary(file_bytes)),
        )
        connection.execute("DELETE FROM welfare_compare_reference_records WHERE import_id = ?", (import_id,))
        connection.executemany(
            """
            INSERT INTO welfare_compare_reference_records
              (id, import_id, row_number, game_name, subtitle, category, name_key, subtitle_key, match_key, raw_json, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    f"{import_id}-{index}",
                    import_id,
                    record["rowNumber"],
                    record["gameName"],
                    record["subtitle"],
                    record["category"],
                    record["nameKey"],
                    record["subtitleKey"],
                    record["matchKey"],
                    json.dumps(record, ensure_ascii=False),
                    imported_at,
                )
                for index, record in enumerate(records, start=1)
            ],
        )
        connection.execute(
            """
            INSERT INTO operation_logs (id, type, target_type, target_id, title, detail, payload_json, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                f"op-welfare-ref-{int(datetime.now(timezone.utc).timestamp() * 1000)}",
                "数据导入",
                "welfare_compare_reference",
                import_id,
                "导入福利对比参考名单",
                f"{xlsx_path.name}：已缓存 {len(records)} 条我方福利游戏记录。",
                json.dumps({"file": str(xlsx_path), "importId": import_id, "rows": len(records), "sha256": file_hash}, ensure_ascii=False),
                imported_at,
            ),
        )
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()
    return {
        "ok": True,
        "importId": import_id,
        "file": str(xlsx_path),
        "sha256": file_hash,
        "sheetName": sheet_name,
        "rowCount": len(records),
        "importedAt": imported_at,
    }


def main():
    parser = argparse.ArgumentParser(description="Import welfare compare reference workbook into SQLite.")
    parser.add_argument("xlsx_path")
    parser.add_argument("--db", required=True)
    args = parser.parse_args()
    print(json.dumps(import_reference(Path(args.db), Path(args.xlsx_path)), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
