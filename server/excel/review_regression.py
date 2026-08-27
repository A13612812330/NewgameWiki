
import json
import re
import sys
from pathlib import Path
from difflib import SequenceMatcher

from openpyxl import load_workbook


def clean(value):
    text = str(value or "")
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def compact(value):
    return re.sub(r"[\s，。！？、,.!?:：《》“”'\"（）()；;：-]", "", clean(value).lower())


def score(current, expected):
    cur = compact(current)
    exp = compact(expected)
    if not cur or not exp:
        return 0.0
    if exp in cur or cur in exp:
        return 1.0 if exp in cur else 0.85
    return round(SequenceMatcher(None, cur, exp).ratio(), 3)


def main():
    xlsx_path = Path(sys.argv[1] if len(sys.argv) > 1 else "data/excel/手游池整理表.xlsx")
    output_path = Path(sys.argv[2] if len(sys.argv) > 2 else "data/reports/review-regression.json")
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb["爬虫字段审查表"]
    headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    idx = {name: pos + 1 for pos, name in enumerate(headers) if name}
    rows = []
    for row in range(2, ws.max_row + 1):
        expected = ws.cell(row, idx["用户指定保留字段"]).value
        if not expected:
            continue
        current = ws.cell(row, idx["整理内容"]).value
        item = {
            "行号": row,
            "游戏名称": ws.cell(row, idx["游戏名称"]).value,
            "来源": ws.cell(row, idx["来源"]).value,
            "审查层级": ws.cell(row, idx["审查层级"]).value,
            "当前整理内容": current or "",
            "用户指定保留字段": expected or "",
            "相似度": score(current, expected),
        }
        if item["相似度"] >= 0.8:
            item["结果"] = "接近"
        elif item["相似度"] >= 0.45:
            item["结果"] = "需复核"
        else:
            item["结果"] = "未命中"
        rows.append(item)
    wb.close()
    summary = {
        "样本数": len(rows),
        "接近": sum(1 for item in rows if item["结果"] == "接近"),
        "需复核": sum(1 for item in rows if item["结果"] == "需复核"),
        "未命中": sum(1 for item in rows if item["结果"] == "未命中"),
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps({"summary": summary, "items": rows}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
