import json
import sys
from pathlib import Path

from openpyxl import load_workbook


NAME_KEYS = {"游戏名称", "游戏名", "名称", "name", "gameName"}
SUBTITLE_KEYS = {"副标题", "折扣标签", "福利标签", "版本", "subtitle", "subTitle"}
CATEGORY_KEYS = {"类别", "分类", "类型", "category"}


def normalize_header(value):
    return str(value or "").strip()


def find_column(headers, candidates):
    normalized = [normalize_header(item) for item in headers]
    for candidate in candidates:
        if candidate in normalized:
            return normalized.index(candidate)
    for index, header in enumerate(normalized):
        if any(candidate in header for candidate in candidates if len(candidate) >= 2):
            return index
    return -1


def main():
    if len(sys.argv) < 2:
        raise SystemExit("usage: parse_welfare_compare.py <xlsx_path>")
    xlsx_path = Path(sys.argv[1])
    if not xlsx_path.exists():
        raise SystemExit(f"文件不存在：{xlsx_path}")
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    records = []
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        try:
            headers = next(rows)
        except StopIteration:
            continue
        name_index = find_column(headers, NAME_KEYS)
        subtitle_index = find_column(headers, SUBTITLE_KEYS)
        category_index = find_column(headers, CATEGORY_KEYS)
        if name_index < 0:
            continue
        for row in rows:
            game_name = str(row[name_index] or "").strip() if name_index < len(row) else ""
            subtitle = str(row[subtitle_index] or "").strip() if 0 <= subtitle_index < len(row) else ""
            category = str(row[category_index] or "").strip() if 0 <= category_index < len(row) else ""
            if not game_name:
                continue
            records.append({
                "gameName": game_name,
                "subtitle": subtitle,
                "category": category,
                "sheet": ws.title,
            })
    wb.close()
    print(json.dumps({"records": records}, ensure_ascii=False))


if __name__ == "__main__":
    main()
