import argparse
import json
import re
import sqlite3
import time
import urllib.error
import urllib.request
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


NAME_COL = 2
SUBTITLE_COL = 3
CATEGORY_COL = 4
DOWNLOAD_COL = 5
LEVEL_RANK = {"C": 0, "B": 1, "A": 2, "S": 3}


def normalize_text(value):
    return re.sub(
        r"[\s\"'“”‘’`~!！?？,，.。:：;；、|/\\\-_=+·・《》<>【】\[\]（）(){}]",
        "",
        str(value or ""),
    ).lower()


def parse_number(value):
    if value is None:
        return 0
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).strip().replace(",", "")
    wan = re.search(r"(\d+(?:\.\d+)?)\s*万", text)
    if wan:
        return int(float(wan.group(1)) * 10000)
    number = re.search(r"\d+(?:\.\d+)?", text)
    return int(float(number.group(0))) if number else 0


def compute_level(game_name, category, download_count):
    is_single = "单机" in str(category or "")
    if not is_single and download_count >= 3000:
        return "S"
    if not is_single and download_count >= 1000:
        return "A"
    if download_count >= 1000:
        return "B"
    return "C"


def read_importance_rows(xlsx_path):
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        next(rows, None)
        records = []
        for index, row in enumerate(rows, start=2):
            game_name = str(row[NAME_COL] or "").strip() if len(row) > NAME_COL else ""
            if not game_name:
                continue
            subtitle = str(row[SUBTITLE_COL] or "").strip() if len(row) > SUBTITLE_COL else ""
            category = str(row[CATEGORY_COL] or "").strip() if len(row) > CATEGORY_COL else ""
            download_count = parse_number(row[DOWNLOAD_COL] if len(row) > DOWNLOAD_COL else 0)
            records.append({
                "row": index,
                "gameName": game_name,
                "subtitle": subtitle,
                "category": category,
                "downloadCount": download_count,
                "importanceLevel": compute_level(game_name, category, download_count),
                "nameKey": normalize_text(game_name),
                "fullKey": f"{normalize_text(game_name)}::{normalize_text(subtitle)}",
            })
        return records
    finally:
        workbook.close()


def subtitle_candidates(payload):
    candidates = []
    for key in ("subtitle", "benefitSubtitle", "welfareTitle", "rawStatus", "status"):
        value = str(payload.get(key) or "").strip()
        if value and value not in candidates:
            candidates.append(value)
    return candidates or [""]


def is_welfare_record(pool, payload):
    text = " ".join([
        pool,
        str(payload.get("platform") or ""),
        str(payload.get("pool") or ""),
        str(payload.get("welfareKind") or ""),
        " ".join(payload.get("sources") or []),
    ])
    return pool == "welfare" or "福利" in text or "小七" in text


def build_excel_indexes(rows):
    by_full = {}
    by_name_blank = {}
    for row in rows:
        current = by_full.get(row["fullKey"])
        if not current or LEVEL_RANK[row["importanceLevel"]] > LEVEL_RANK[current["importanceLevel"]]:
            by_full[row["fullKey"]] = row
        if not row["subtitle"]:
            current_blank = by_name_blank.get(row["nameKey"])
            if not current_blank or LEVEL_RANK[row["importanceLevel"]] > LEVEL_RANK[current_blank["importanceLevel"]]:
                by_name_blank[row["nameKey"]] = row
    return by_full, by_name_blank


def match_records(db_path, excel_rows):
    by_full, by_name_blank = build_excel_indexes(excel_rows)
    connection = sqlite3.connect(db_path)
    try:
        rows = connection.execute(
            "SELECT pool, item_id, payload_json FROM game_records WHERE pool IN ('mobile', 'welfare')"
        ).fetchall()
    finally:
        connection.close()

    matched = []
    unmatched = []
    for pool, item_id, payload_json in rows:
        try:
            payload = json.loads(payload_json)
        except json.JSONDecodeError:
            continue
        game_name = str(payload.get("gameName") or "").strip()
        name_key = normalize_text(game_name)
        welfare = is_welfare_record(pool, payload)
        hit = None
        hit_rule = ""
        for subtitle in subtitle_candidates(payload):
            hit = by_full.get(f"{name_key}::{normalize_text(subtitle)}")
            if hit:
                hit_rule = "名称+副标题"
                break
        if not hit and not welfare:
            hit = by_name_blank.get(name_key)
            if hit:
                hit_rule = "名称匹配(名单副标题为空)"
        if not hit:
            unmatched.append({
                "pool": pool,
                "itemId": item_id,
                "gameName": game_name,
                "subtitle": subtitle_candidates(payload)[0],
            })
            continue
        before = str(payload.get("importanceLevel") or "C").upper()
        if before not in LEVEL_RANK:
            before = "C"
        after = hit["importanceLevel"]
        matched.append({
            "pool": pool,
            "itemId": item_id,
            "gameName": game_name,
            "subtitle": subtitle_candidates(payload)[0],
            "before": before,
            "after": after,
            "changed": before != after,
            "downloadCount": hit["downloadCount"],
            "category": hit["category"],
            "excelRow": hit["row"],
            "matchRule": hit_rule,
        })
    return matched, unmatched


def post_json(api_base, route, payload):
    data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    request = urllib.request.Request(
        api_base.rstrip("/") + route,
        data=data,
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    with urllib.request.urlopen(request, timeout=30) as response:
        return json.loads(response.read().decode("utf-8") or "{}")


def apply_via_api(api_base, matches):
    changed = [item for item in matches if item["changed"]]
    if not changed:
        return {"requested": 0, "updated": 0, "failed": []}
    post_json(api_base, "/api/backup/create", {})
    updated = 0
    failed = []
    for item in changed:
        try:
            post_json(api_base, "/api/pool/field-choice", {
                "gameId": item["itemId"],
                "field": "importanceLevel",
                "value": item["after"],
                "source": f"游戏名单导入：{item['category']} / 七天下载 {item['downloadCount']}",
            })
            updated += 1
            time.sleep(0.03)
        except (urllib.error.URLError, TimeoutError, json.JSONDecodeError) as error:
            failed.append({**item, "error": str(error)})
    return {"requested": len(changed), "updated": updated, "failed": failed}


def main():
    parser = argparse.ArgumentParser(description="Apply importance levels from a game management workbook.")
    parser.add_argument("xlsx_path")
    parser.add_argument("--db", required=True)
    parser.add_argument("--api", default="http://127.0.0.1:64111")
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--output", default="")
    args = parser.parse_args()

    xlsx_path = Path(args.xlsx_path)
    db_path = Path(args.db)
    excel_rows = read_importance_rows(xlsx_path)
    matched, unmatched = match_records(db_path, excel_rows)
    apply_result = apply_via_api(args.api, matched) if args.apply else None
    result = {
        "ok": True,
        "xlsx": str(xlsx_path),
        "db": str(db_path),
        "excelRows": len(excel_rows),
        "matched": len(matched),
        "changed": sum(1 for item in matched if item["changed"]),
        "unmatchedDb": len(unmatched),
        "levelCounts": dict(Counter(item["after"] for item in matched)),
        "ruleCounts": dict(Counter(item["matchRule"] for item in matched)),
        "apply": apply_result,
        "sampleChanges": [item for item in matched if item["changed"]][:30],
        "sampleUnmatched": unmatched[:30],
    }
    if args.output:
        Path(args.output).parent.mkdir(parents=True, exist_ok=True)
        Path(args.output).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
