import argparse
import hashlib
import json
import re
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


def now_iso():
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def normalize_text(value):
    return re.sub(
        r"[\s\"'“”‘’`~!！?？,，.。:：;；、|/\\\-_=+·・《》<>【】\[\]（）(){}]",
        "",
        str(value or ""),
    ).lower()


def parse_int(value):
    if value is None:
        return 0
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).strip().replace(",", "")
    wan = re.search(r"(\d+(?:\.\d+)?)\s*万", text)
    if wan:
        return int(float(wan.group(1)) * 10000)
    number = re.search(r"\d+(?:\.\d+)?", text)
    return int(float(number.group(0))) if number else 0


def header_map(headers):
    return {str(value or "").strip(): index for index, value in enumerate(headers)}


def pick(row, columns, name, default=""):
    index = columns.get(name)
    if index is None or index >= len(row):
        return default
    value = row[index]
    if value is None:
        return default
    return str(value).strip()


def parse_workbook(xlsx_path):
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            return sheet.title, []
        columns = header_map(headers)
        records = []
        for row_number, row in enumerate(rows, start=2):
            game_name = pick(row, columns, "游戏名称")
            if not game_name:
                continue
            subtitle = pick(row, columns, "副标题")
            name_key = normalize_text(game_name)
            subtitle_key = normalize_text(subtitle)
            record = {
                "rowNumber": row_number,
                "platformGameId": pick(row, columns, "游戏ID"),
                "gameName": game_name,
                "subtitle": subtitle,
                "category": pick(row, columns, "类别"),
                "downloads7d": parse_int(row[columns["七天下载量"]]) if "七天下载量" in columns else 0,
                "normalDownloads7d": parse_int(row[columns["七天下载量-常规下载"]]) if "七天下载量-常规下载" in columns else 0,
                "cloudDownloads7d": parse_int(row[columns["七天下载量-畅玩下载"]]) if "七天下载量-畅玩下载" in columns else 0,
                "downloadStatus": pick(row, columns, "下载状态"),
                "nameKey": name_key,
                "subtitleKey": subtitle_key,
                "matchKey": f"{name_key}::{subtitle_key}",
            }
            records.append(record)
        return sheet.title, records
    finally:
        workbook.close()


def ensure_schema(connection):
    connection.executescript(
        """
        CREATE TABLE IF NOT EXISTS platform_game_reference_imports (
          id TEXT PRIMARY KEY,
          file_name TEXT NOT NULL,
          source_path TEXT,
          file_sha256 TEXT NOT NULL UNIQUE,
          file_size INTEGER NOT NULL DEFAULT 0,
          sheet_name TEXT,
          row_count INTEGER NOT NULL DEFAULT 0,
          imported_at TEXT NOT NULL,
          is_active INTEGER NOT NULL DEFAULT 1,
          file_blob BLOB
        );
        CREATE TABLE IF NOT EXISTS platform_game_reference_records (
          id TEXT PRIMARY KEY,
          import_id TEXT NOT NULL,
          row_number INTEGER NOT NULL,
          platform_game_id TEXT,
          game_name TEXT NOT NULL,
          subtitle TEXT,
          category TEXT,
          downloads_7d INTEGER NOT NULL DEFAULT 0,
          normal_downloads_7d INTEGER NOT NULL DEFAULT 0,
          cloud_downloads_7d INTEGER NOT NULL DEFAULT 0,
          download_status TEXT,
          name_key TEXT NOT NULL,
          subtitle_key TEXT NOT NULL DEFAULT '',
          match_key TEXT NOT NULL,
          raw_json TEXT DEFAULT '{}',
          created_at TEXT NOT NULL,
          FOREIGN KEY(import_id) REFERENCES platform_game_reference_imports(id)
        );
        CREATE TABLE IF NOT EXISTS operation_logs (
          id TEXT PRIMARY KEY,
          type TEXT NOT NULL,
          target_type TEXT,
          target_id TEXT,
          title TEXT,
          detail TEXT,
          payload_json TEXT DEFAULT '{}',
          created_at TEXT NOT NULL
        );
        CREATE INDEX IF NOT EXISTS idx_platform_game_reference_active ON platform_game_reference_imports(is_active, imported_at);
        CREATE INDEX IF NOT EXISTS idx_platform_game_reference_match ON platform_game_reference_records(match_key);
        CREATE INDEX IF NOT EXISTS idx_platform_game_reference_name ON platform_game_reference_records(name_key);
        """
    )


def import_reference(db_path, xlsx_path):
    file_bytes = xlsx_path.read_bytes()
    file_hash = hashlib.sha256(file_bytes).hexdigest()
    import_id = f"platform-ref-{file_hash[:16]}"
    imported_at = now_iso()
    sheet_name, records = parse_workbook(xlsx_path)
    connection = sqlite3.connect(db_path)
    try:
        ensure_schema(connection)
        connection.execute("BEGIN")
        connection.execute("UPDATE platform_game_reference_imports SET is_active = 0")
        connection.execute(
            """
            INSERT INTO platform_game_reference_imports
              (id, file_name, source_path, file_sha256, file_size, sheet_name, row_count, imported_at, is_active, file_blob)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, 1, ?)
            ON CONFLICT(file_sha256) DO UPDATE SET
              file_name = excluded.file_name,
              source_path = excluded.source_path,
              file_size = excluded.file_size,
              sheet_name = excluded.sheet_name,
              row_count = excluded.row_count,
              imported_at = excluded.imported_at,
              is_active = 1,
              file_blob = excluded.file_blob
            """,
            (
                import_id,
                xlsx_path.name,
                str(xlsx_path),
                file_hash,
                len(file_bytes),
                sheet_name,
                len(records),
                imported_at,
                sqlite3.Binary(file_bytes),
            ),
        )
        connection.execute("DELETE FROM platform_game_reference_records WHERE import_id = ?", (import_id,))
        connection.executemany(
            """
            INSERT INTO platform_game_reference_records
              (id, import_id, row_number, platform_game_id, game_name, subtitle, category,
               downloads_7d, normal_downloads_7d, cloud_downloads_7d, download_status,
               name_key, subtitle_key, match_key, raw_json, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    f"{import_id}-{record['rowNumber']}",
                    import_id,
                    record["rowNumber"],
                    record["platformGameId"],
                    record["gameName"],
                    record["subtitle"],
                    record["category"],
                    record["downloads7d"],
                    record["normalDownloads7d"],
                    record["cloudDownloads7d"],
                    record["downloadStatus"],
                    record["nameKey"],
                    record["subtitleKey"],
                    record["matchKey"],
                    json.dumps(record, ensure_ascii=False),
                    imported_at,
                )
                for record in records
            ],
        )
        connection.execute(
            """
            INSERT INTO operation_logs (id, type, target_type, target_id, title, detail, payload_json, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                f"op-platform-ref-{int(datetime.now(timezone.utc).timestamp() * 1000)}",
                "数据导入",
                "platform_game_reference",
                import_id,
                "导入我方平台游戏参考名单",
                f"{xlsx_path.name}：已保存 {len(records)} 条平台游戏记录。",
                json.dumps({"file": str(xlsx_path), "importId": import_id, "rows": len(records), "sha256": file_hash}, ensure_ascii=False),
                imported_at,
            ),
        )
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()
    return {
        "ok": True,
        "importId": import_id,
        "file": str(xlsx_path),
        "sha256": file_hash,
        "sheetName": sheet_name,
        "rowCount": len(records),
        "importedAt": imported_at,
    }


def main():
    parser = argparse.ArgumentParser(description="Import platform game list into SQLite reference tables.")
    parser.add_argument("xlsx_path")
    parser.add_argument("--db", required=True)
    args = parser.parse_args()
    result = import_reference(Path(args.db), Path(args.xlsx_path))
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
